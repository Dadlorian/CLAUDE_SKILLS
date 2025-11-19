#!/usr/bin/env python3
"""
Basic document generation example using Python-docx
Demonstrates simple template variable substitution in Word documents
"""

from docx import Document
from docx.shared import Pt, RGBColor
from datetime import datetime

class BasicDocumentGenerator:
    def __init__(self, template_path):
        """Initialize with a template document."""
        self.template = Document(template_path)
    
    def replace_text_in_paragraphs(self, replacements):
        """Replace text in all paragraphs."""
        for paragraph in self.template.paragraphs:
            for key, value in replacements.items():
                if key in paragraph.text:
                    # Replace text while preserving formatting
                    for run in paragraph.runs:
                        if key in run.text:
                            run.text = run.text.replace(key, value)
    
    def replace_text_in_tables(self, replacements):
        """Replace text in tables."""
        for table in self.template.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for key, value in replacements.items():
                            if key in paragraph.text:
                                paragraph.text = paragraph.text.replace(key, value)
    
    def generate(self, data, output_path):
        """Generate document with provided data."""
        replacements = {
            '{CLIENT_NAME}': data.get('client_name', ''),
            '{CLIENT_ADDRESS}': data.get('client_address', ''),
            '{DATE}': data.get('date', datetime.now().strftime('%B %d, %Y')),
            '{AMOUNT}': data.get('amount', ''),
            '{MATTER_ID}': data.get('matter_id', ''),
        }
        
        self.replace_text_in_paragraphs(replacements)
        self.replace_text_in_tables(replacements)
        
        self.template.save(output_path)
        return output_path


# Example usage
if __name__ == '__main__':
    data = {
        'client_name': 'John Doe',
        'client_address': '123 Main Street, New York, NY 10001',
        'amount': '$5,000.00',
        'matter_id': 'MATTER-2024-001'
    }
    
    generator = BasicDocumentGenerator('templates/engagement_letter.docx')
    output = generator.generate(data, 'output/engagement_letter_final.docx')
    print(f"Document generated: {output}")
