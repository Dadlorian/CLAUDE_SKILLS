#!/usr/bin/env python3
"""
Basic DOCX Generation with python-docx
Demonstrates creating a simple legal document from scratch
"""

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from datetime import datetime

def create_basic_agreement(client_name, effective_date, service_description):
    """Generate a basic service agreement"""

    # Create document
    doc = Document()

    # Add title
    title = doc.add_heading('SERVICE AGREEMENT', 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    # Add effective date
    date_para = doc.add_paragraph(f'Effective Date: {effective_date}')
    date_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    # Add parties section
    doc.add_heading('PARTIES', level=1)
    doc.add_paragraph(
        f'This Agreement is entered into between {client_name} ("Client") '
        'and Law Firm Professional Corporation ("Firm").'
    )

    # Add services section
    doc.add_heading('SERVICES', level=1)
    doc.add_paragraph(
        f'The Firm shall provide the following legal services: {service_description}'
    )

    # Add fees section
    doc.add_heading('FEES', level=1)
    doc.add_paragraph(
        'Client agrees to pay fees as set forth in Schedule A attached hereto.'
    )

    # Add signature blocks
    doc.add_page_break()
    doc.add_heading('SIGNATURES', level=1)

    # Client signature
    doc.add_paragraph('CLIENT:')
    doc.add_paragraph()
    doc.add_paragraph('_' * 50)
    doc.add_paragraph(client_name)
    doc.add_paragraph('Date: _______________')
    doc.add_paragraph()

    # Firm signature
    doc.add_paragraph('FIRM:')
    doc.add_paragraph()
    doc.add_paragraph('_' * 50)
    doc.add_paragraph('Law Firm Professional Corporation')
    doc.add_paragraph('Date: _______________')

    # Save document
    filename = f'Service_Agreement_{client_name.replace(" ", "_")}.docx'
    doc.save(filename)
    print(f'Created: {filename}')

    return filename

if __name__ == '__main__':
    # Example usage
    create_basic_agreement(
        client_name='Acme Corporation',
        effective_date=datetime.now().strftime('%B %d, %Y'),
        service_description='Corporate formation and governance advice'
    )
