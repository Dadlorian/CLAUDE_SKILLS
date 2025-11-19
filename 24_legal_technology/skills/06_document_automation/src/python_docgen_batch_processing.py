#!/usr/bin/env python3
"""
Batch processing for document generation
Process multiple records and generate documents in parallel
"""

import concurrent.futures
import csv
from pathlib import Path
from docx import Document
from datetime import datetime

class BatchDocumentGenerator:
    def __init__(self, template_path):
        self.template_path = template_path
    
    def generate_single(self, record, output_dir):
        """Generate a single document from a record."""
        try:
            doc = Document(self.template_path)
            
            # Replace placeholders
            for paragraph in doc.paragraphs:
                for key, value in record.items():
                    placeholder = '{' + key.upper() + '}'
                    if placeholder in paragraph.text:
                        paragraph.text = paragraph.text.replace(placeholder, str(value))
            
            # Create unique filename
            filename = f"{record.get('matter_id', 'document')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
            output_path = Path(output_dir) / filename
            
            doc.save(output_path)
            return {'status': 'success', 'file': output_path, 'record_id': record.get('matter_id')}
        
        except Exception as e:
            return {'status': 'error', 'error': str(e), 'record_id': record.get('matter_id')}
    
    def batch_generate(self, csv_file, output_dir, max_workers=4):
        """Generate multiple documents from CSV file."""
        Path(output_dir).mkdir(parents=True, exist_ok=True)
        
        records = []
        with open(csv_file, 'r') as f:
            reader = csv.DictReader(f)
            records = list(reader)
        
        results = []
        with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as executor:
            futures = [executor.submit(self.generate_single, record, output_dir) 
                      for record in records]
            
            for future in concurrent.futures.as_completed(futures):
                result = future.result()
                results.append(result)
        
        return self._summarize_results(results)
    
    def _summarize_results(self, results):
        """Summarize batch generation results."""
        successful = [r for r in results if r['status'] == 'success']
        failed = [r for r in results if r['status'] == 'error']
        
        return {
            'total': len(results),
            'successful': len(successful),
            'failed': len(failed),
            'details': {
                'successful': successful,
                'failed': failed
            }
        }


# Example CSV data structure
# matter_id,client_name,client_address,date,amount
# MATTER-001,Client A,123 Main St,2024-01-15,$10000
# MATTER-002,Client B,456 Oak Ave,2024-01-15,$15000

# Example usage
if __name__ == '__main__':
    generator = BatchDocumentGenerator('templates/fee_agreement.docx')
    
    results = generator.batch_generate(
        'data/clients.csv',
        'output/generated_documents',
        max_workers=4
    )
    
    print(f"Batch Generation Summary:")
    print(f"Total: {results['total']}")
    print(f"Successful: {results['successful']}")
    print(f"Failed: {results['failed']}")
