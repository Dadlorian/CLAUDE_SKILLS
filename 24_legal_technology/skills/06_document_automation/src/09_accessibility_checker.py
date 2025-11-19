#!/usr/bin/env python3
"""Accessibility Checker for Generated Documents"""

from docx import Document
from PyPDF2 import PdfReader
import re

class AccessibilityChecker:
    def check_docx(self, file_path):
        """Check DOCX for accessibility issues"""
        issues = []
        doc = Document(file_path)

        # Check for heading structure
        headings = [p for p in doc.paragraphs if p.style.name.startswith('Heading')]
        if not headings:
            issues.append("No headings found - add document structure")

        # Check for alt text on images
        # (Simplified - full implementation would check all images)
        if any('image' in str(rel.target_ref) for rel in doc.part.rels.values()):
            issues.append("Verify alt text for all images")

        # Check for proper list usage
        for para in doc.paragraphs:
            text = para.text.strip()
            if text and text[0].isdigit() and '. ' in text[:5]:
                if para.style.name == 'Normal':
                    issues.append(f"Use list style instead of manual numbering: {text[:50]}")

        return issues

    def check_pdf(self, file_path):
        """Check PDF for accessibility issues"""
        issues = []
        reader = PdfReader(file_path)

        # Check for searchable text
        text = ''
        for page in reader.pages:
            text += page.extract_text()

        if len(text.strip()) < 100:
            issues.append("PDF appears to be image-only or has minimal text")

        # Check metadata
        metadata = reader.metadata
        if not metadata or not metadata.get('/Title'):
            issues.append("PDF missing title metadata")

        # Check for tags (simplified check)
        if '/MarkInfo' not in reader.trailer.get('/Root', {}):
            issues.append("PDF may not be tagged for accessibility")

        return issues

    def generate_report(self, file_path):
        """Generate accessibility report"""
        if file_path.endswith('.docx'):
            issues = self.check_docx(file_path)
        elif file_path.endswith('.pdf'):
            issues = self.check_pdf(file_path)
        else:
            return ["Unsupported file format"]

        return {
            'file': file_path,
            'total_issues': len(issues),
            'issues': issues,
            'passed': len(issues) == 0
        }

if __name__ == '__main__':
    checker = AccessibilityChecker()
    # Example usage (requires existing file)
    # report = checker.generate_report('agreement.pdf')
    # print(json.dumps(report, indent=2))
    print("Accessibility checker ready")
