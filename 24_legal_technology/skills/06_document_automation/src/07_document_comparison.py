#!/usr/bin/env python3
"""Document Comparison and Redlining"""

from docx import Document
import difflib

def compare_documents(file1, file2):
    """Compare two documents and show differences"""

    # Read documents
    doc1 = Document(file1)
    doc2 = Document(file2)

    # Extract text
    text1 = '\n'.join([para.text for para in doc1.paragraphs])
    text2 = '\n'.join([para.text for para in doc2.paragraphs])

    # Generate diff
    diff = difflib.unified_diff(
        text1.splitlines(),
        text2.splitlines(),
        fromfile=file1,
        tofile=file2,
        lineterm=''
    )

    # Output differences
    differences = list(diff)
    if differences:
        print("Differences found:")
        for line in differences:
            print(line)
    else:
        print("Documents are identical")

    return differences

def generate_redline_html(file1, file2, output='redline.html'):
    """Generate HTML redline comparison"""

    doc1 = Document(file1)
    doc2 = Document(file2)

    text1 = [para.text for para in doc1.paragraphs]
    text2 = [para.text for para in doc2.paragraphs]

    diff = difflib.HtmlDiff()
    html = diff.make_file(text1, text2, fromdesc=file1, todesc=file2)

    with open(output, 'w') as f:
        f.write(html)

    print(f'Created redline: {output}')
    return output

if __name__ == '__main__':
    # Example usage (requires two existing DOCX files)
    # compare_documents('version1.docx', 'version2.docx')
    # generate_redline_html('version1.docx', 'version2.docx')
    print("Document comparison module ready")
