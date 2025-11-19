"""
DOCX Document Generator
Generates Microsoft Word documents programmatically for legal contracts, agreements, and forms.

Dependencies: python-docx
Install: pip install python-docx
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime
from typing import Dict, List, Optional


class DocxGenerator:
    """Generate formatted DOCX documents for legal purposes."""

    def __init__(self, title: str = "Legal Document"):
        """Initialize document with title."""
        self.doc = Document()
        self.title = title
        self._add_title(title)

    def _add_title(self, title: str):
        """Add formatted title to document."""
        title_para = self.doc.add_paragraph(title)
        title_para.style = 'Heading 1'
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Format title
        for run in title_para.runs:
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 51, 102)  # Dark blue

    def add_heading(self, text: str, level: int = 1):
        """Add heading at specified level."""
        heading = self.doc.add_paragraph(text)
        heading.style = f'Heading {level}'
        return heading

    def add_paragraph(self, text: str, bold: bool = False, italic: bool = False) -> None:
        """Add formatted paragraph."""
        para = self.doc.add_paragraph(text)
        for run in para.runs:
            run.font.bold = bold
            run.font.italic = italic
            run.font.size = Pt(11)

    def add_signature_block(self, party_name: str):
        """Add signature block for contract signing."""
        self.doc.add_paragraph()  # Blank line
        self.doc.add_paragraph("_" * 50)
        self.doc.add_paragraph(party_name, style='Normal')

    def add_table(self, rows: int, cols: int, data: List[List[str]]) -> None:
        """Add table with data."""
        table = self.doc.add_table(rows=rows, cols=cols)
        table.style = 'Light Grid Accent 1'

        for i, row_data in enumerate(data):
            if i < len(table.rows):
                for j, cell_data in enumerate(row_data):
                    if j < len(table.rows[i].cells):
                        table.rows[i].cells[j].text = str(cell_data)

    def add_numbered_list(self, items: List[str]) -> None:
        """Add numbered list."""
        for item in items:
            para = self.doc.add_paragraph(item, style='List Number')

    def add_bulleted_list(self, items: List[str]) -> None:
        """Add bulleted list."""
        for item in items:
            para = self.doc.add_paragraph(item, style='List Bullet')

    def add_footer_with_date(self) -> None:
        """Add footer with current date."""
        section = self.doc.sections[0]
        footer = section.footer
        footer_para = footer.paragraphs[0]
        footer_para.text = f"Document generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"

    def save(self, filename: str) -> None:
        """Save document to file."""
        self.doc.save(filename)
        print(f"Document saved to {filename}")


def example_service_agreement():
    """Generate example service agreement."""
    doc = DocxGenerator("SERVICE AGREEMENT")

    doc.add_heading("1. PARTIES", level=2)
    doc.add_paragraph("This Agreement is between Client, LLC ('Client') and Service Provider, Inc. ('Provider').")

    doc.add_heading("2. SERVICES", level=2)
    doc.add_paragraph("Provider agrees to provide the following services:")
    services = [
        "Legal document review and analysis",
        "Contract drafting and negotiation",
        "Compliance consulting",
        "Legal research and memoranda"
    ]
    doc.add_numbered_list(services)

    doc.add_heading("3. FEES", level=2)
    table_data = [
        ["Service Type", "Rate", "Unit"],
        ["Document Review", "$250", "Hour"],
        ["Drafting", "$300", "Hour"],
        ["Consulting", "$275", "Hour"]
    ]
    doc.add_table(len(table_data), len(table_data[0]), table_data)

    doc.add_heading("4. TERM", level=2)
    doc.add_paragraph("This Agreement shall commence on the date signed and continue for one (1) year.")

    doc.add_heading("5. CONFIDENTIALITY", level=2)
    doc.add_paragraph("Both parties agree to maintain strict confidentiality of all proprietary information.")

    doc.add_paragraph()
    doc.add_signature_block("Client Signature: ___________________  Date: __________")
    doc.add_paragraph()
    doc.add_signature_block("Provider Signature: ___________________  Date: __________")

    doc.add_footer_with_date()
    return doc


def example_nda():
    """Generate example Non-Disclosure Agreement."""
    doc = DocxGenerator("NON-DISCLOSURE AGREEMENT")

    doc.add_heading("1. CONFIDENTIAL INFORMATION", level=2)
    doc.add_paragraph(
        "Confidential Information means all non-public information disclosed by one party "
        "to the other, including but not limited to trade secrets, business plans, financial data, "
        "technical data, and client lists."
    )

    doc.add_heading("2. OBLIGATIONS", level=2)
    obligations = [
        "Maintain confidentiality using reasonable care",
        "Limit disclosure to employees with need-to-know",
        "Protect information for a period of three (3) years",
        "Return or destroy information upon request"
    ]
    doc.add_numbered_list(obligations)

    doc.add_heading("3. EXCLUSIONS", level=2)
    exclusions = [
        "Information already public",
        "Information rightfully obtained from third parties",
        "Information independently developed",
        "Information required to be disclosed by law"
    ]
    doc.add_bulleted_list(exclusions)

    doc.add_paragraph()
    doc.add_signature_block("Disclosing Party: ___________________  Date: __________")
    doc.add_paragraph()
    doc.add_signature_block("Receiving Party: ___________________  Date: __________")

    doc.add_footer_with_date()
    return doc


if __name__ == "__main__":
    # Generate example documents
    service_agreement = example_service_agreement()
    service_agreement.save("/tmp/service_agreement.docx")

    nda = example_nda()
    nda.save("/tmp/nda.docx")

    print("\nExample documents generated successfully!")
