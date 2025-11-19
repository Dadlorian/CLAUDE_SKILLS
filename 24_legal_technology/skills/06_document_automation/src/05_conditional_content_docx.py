#!/usr/bin/env python3
"""Conditional Content Generation in DOCX"""

from docx import Document
from docx.shared import Pt

def generate_contract_with_conditionals(data):
    doc = Document()
    doc.add_heading('PURCHASE AGREEMENT', 0)

    # Basic provisions
    doc.add_paragraph(f"Buyer: {data['buyer_name']}")
    doc.add_paragraph(f"Seller: {data['seller_name']}")
    doc.add_paragraph(f"Purchase Price: ${data['purchase_price']:,.2f}")

    # Conditional: Earnout
    if data.get('include_earnout'):
        doc.add_heading('EARNOUT PROVISIONS', level=1)
        doc.add_paragraph(
            f"Maximum Earnout Amount: ${data['earnout_amount']:,.2f}"
        )
        doc.add_paragraph(f"Earnout Period: {data['earnout_period']} years")

    # Conditional: Escrow
    if data.get('include_escrow'):
        doc.add_heading('ESCROW', level=1)
        doc.add_paragraph(
            f"Escrow Amount: ${data['escrow_amount']:,.2f}"
        )
        doc.add_paragraph(f"Escrow Period: {data['escrow_period']} months")

    # Conditional: Non-Compete
    if data.get('include_noncompete'):
        doc.add_heading('NON-COMPETE', level=1)
        doc.add_paragraph(
            f"Duration: {data['noncompete_duration']} years"
        )
        doc.add_paragraph(
            f"Territory: {data['noncompete_territory']}"
        )

    filename = f"Agreement_{data['buyer_name'].replace(' ', '_')}.docx"
    doc.save(filename)
    print(f'Created: {filename}')
    return filename

if __name__ == '__main__':
    data = {
        'buyer_name': 'Acme Corp',
        'seller_name': 'Smith Industries',
        'purchase_price': 5000000,
        'include_earnout': True,
        'earnout_amount': 1000000,
        'earnout_period': 3,
        'include_escrow': True,
        'escrow_amount': 500000,
        'escrow_period': 12,
        'include_noncompete': True,
        'noncompete_duration': 2,
        'noncompete_territory': 'California'
    }
    generate_contract_with_conditionals(data)
